from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PRODUCTS_SHEET = "Products"
ALIASES_SHEET = "Product_Aliases"


PRODUCT_REQUIRED_HEADERS = {
    "Product ID",
    "Normalized Product",
    "Category",
    "Subcategory",
    "Default Brand",
    "Comparison Unit",
    "Active",
    "Notes",
}


ALIAS_REQUIRED_HEADERS = {
    "Receipt Text / Alias",
    "Product ID",
    "Store",
    "Status",
    "Notes",
}


def _build_header_map(worksheet) -> dict[str, int]:
    """
    Build a mapping from header name to Excel column index, instead of relying on hardcoded column indices.
    """

    header_map = {}

    for cell in worksheet[1]:
        if cell.value is None:
            continue

        header = str(cell.value).strip() # .strip() to remove leading/trailing whitespace
        header_map[header] = cell.column

    return header_map


def _validate_headers(
    header_map: dict[str, int],
    required_headers: set[str],
    sheet_name: str,
) -> None:
    """
    Validate that a worksheet contains all required headers.
    """

    missing_headers = required_headers - header_map.keys() #.keys() returns a view of the dict's keys, can be treated like a set.

    if missing_headers:
        missing = ", ".join(sorted(missing_headers))

        raise ValueError(
            f"Sheet '{sheet_name}' is missing required headers: {missing}"
        )


def load_products(excel_path: str | Path) -> list[dict[str, Any]]:
    """
    Load active products from the Products worksheet.
    """

    workbook = load_workbook( 
        filename=excel_path,
        data_only=True,
        read_only=True,
    ) # load_workbook opens the Excel file in read-only mode, data_only ensures we get the values rather than the formulas themselves.

    try:
        if PRODUCTS_SHEET not in workbook.sheetnames:
            raise ValueError(
                f"Workbook does not contain the '{PRODUCTS_SHEET}' sheet."
            )

        worksheet = workbook[PRODUCTS_SHEET]

        header_map = _build_header_map(worksheet)

        _validate_headers(
            header_map=header_map,
            required_headers=PRODUCT_REQUIRED_HEADERS,
            sheet_name=PRODUCTS_SHEET,
        )

        products = []

        for row in worksheet.iter_rows(
            min_row=2,  # min_row=2 skips the header row
            values_only=True,
        ):
            product_id = row[header_map["Product ID"] - 1]

            if product_id is None:
                continue

            active = row[header_map["Active"] - 1]

            if str(active).strip().casefold() != "yes":
                continue

            product = {
                "product_id": product_id,
                "normalized_name": row[header_map["Normalized Product"] - 1],
                "category": row[header_map["Category"] - 1],
                "subcategory": row[header_map["Subcategory"] - 1],
                "default_brand": row[header_map["Default Brand"] - 1],
                "comparison_unit": row[header_map["Comparison Unit"] - 1],
                "notes": row[header_map["Notes"] - 1],
            }

            products.append(product)

        return products

    finally:
        workbook.close()


def load_aliases(excel_path: str | Path,) -> list[dict[str, Any]]:
    """
    Load confirmed product aliases from the Product_Aliases worksheet.
    """

    workbook = load_workbook(
        filename=excel_path,
        data_only=True,
        read_only=True,
    )

    try:
        if ALIASES_SHEET not in workbook.sheetnames:
            raise ValueError(
                f"Workbook does not contain the '{ALIASES_SHEET}' sheet."
            )

        worksheet = workbook[ALIASES_SHEET]

        header_map = _build_header_map(worksheet)

        _validate_headers(
            header_map=header_map,
            required_headers=ALIAS_REQUIRED_HEADERS,
            sheet_name=ALIASES_SHEET,
        )

        aliases = []

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):
            receipt_text = row[header_map["Receipt Text / Alias"] - 1]

            if receipt_text is None:
                continue

            status = row[header_map["Status"] - 1]

            if str(status).strip().casefold() != "confirmed":
                continue

            alias = {
                "receipt_text": receipt_text,
                "product_id": row[header_map["Product ID"] - 1],
                "store": row[header_map["Store"] - 1],
            }

            aliases.append(alias)

        return aliases

    finally:
        workbook.close()