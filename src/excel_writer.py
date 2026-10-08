from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.receipt_schema import ReceiptExtraction


RECEIPTS_SHEET = "Receipts"
PRICE_HISTORY_SHEET = "Price_History"
REVIEW_SHEET = "Review"


RECEIPTS_REQUIRED_HEADERS = {
    "Receipt ID",
    "Date",
    "Store",
    "Gross Total (€)",
    "Discounts (€)",
    "Total Paid (€)",
    "Line Count",
    "Payment Method",
    "Notes",
}


PRICE_HISTORY_REQUIRED_HEADERS = {
    "Record ID",
    "Receipt ID",
    "Date",
    "Store",
    "Product ID",
    "Receipt/Original Product",
    "Original Quantity",
    "Comparable Quantity",
    "Comparable Unit",
    "Gross Price (€)",
    "Discount (€)",
    "Price Paid (€)",
    "Comparable Price",
    "Notes",
}


REVIEW_REQUIRED_HEADERS = {
    "Review ID",
    "Receipt ID",
    "Original Text",
    "Issue Type",
    "AI Suggestion",
    "Confidence",
    "Status",
    "Correction",
}


def _build_header_map(worksheet) -> dict[str, int]:
    """
    Build a mapping from header name to Excel column index.
    """

    header_map = {}

    for cell in worksheet[1]:
        if cell.value is None:
            continue

        header = str(cell.value).strip()

        header_map[header] = cell.column

    return header_map


def _validate_headers(
    header_map: dict[str, int],
    required_headers: set[str],
    sheet_name: str,
) -> None:
    """
    Validate that all required headers exist in a worksheet.
    """

    missing_headers = required_headers - header_map.keys()

    if missing_headers:
        missing = ", ".join(sorted(missing_headers))

        raise ValueError(
            f"Sheet '{sheet_name}' is missing required headers: "
            f"{missing}"
        )


def _find_next_data_row(worksheet, id_column_index: int) -> int:
    """
    Find the first available row after the last non-empty ID.
    """

    last_data_row = 1

    for row_number in range(2, worksheet.max_row + 1):
        value = worksheet.cell(row=row_number, column=id_column_index).value

        if value is not None and str(value).strip():
            last_data_row = row_number

    return last_data_row + 1


def _write_row(worksheet, header_map: dict[str, int], values: dict[str, Any], id_header: str) -> int:
    """
    Write values to the first available data row. To avoid using worksheet.append when adding a new row, since there was a formatting issue.

    Returns the row number that was written.
    """

    if id_header not in header_map:
        raise ValueError(
            f"ID header '{id_header}' was not found "
            f"in sheet '{worksheet.title}'."
        )

    row_number = _find_next_data_row(
        worksheet=worksheet,
        id_column_index=header_map[id_header],
    )

    for header, value in values.items():

        if header not in header_map:
            raise ValueError(
                f"Cannot write unknown header: '{header}'"
            )

        worksheet.cell(
            row=row_number,
            column=header_map[header],
            value=value,
        )

    return row_number


def _parse_receipt_date(
    value: str | None,
):
    """
    Convert a YYYY-MM-DD receipt date into a Python date object.
    """

    if value is None:
        return None

    try:
        return datetime.strptime(
            value,
            "%Y-%m-%d",
        ).date()

    except ValueError as error:
        raise ValueError(
            f"Invalid receipt date: '{value}'. "
            "Expected YYYY-MM-DD."
        ) from error


def _build_receipt_notes(
    result: ReceiptExtraction,
) -> str | None:
    """
    Build receipt-level notes from branch information and warnings.
    """

    note_parts = []

    if result.ticket.store_branch:
        note_parts.append(f"Branch: {result.ticket.store_branch}")

    if result.warnings:
        warnings_text = " | ".join(result.warnings)

        note_parts.append(f"Warnings: {warnings_text}")

    if not note_parts:
        return None

    return " | ".join(note_parts)


def _append_receipt(
    worksheet,
    header_map: dict[str, int],
    receipt_id: str,
    result: ReceiptExtraction,
) -> None:
    """
    Append one receipt-level row to the Receipts worksheet.
    """

    values = {
        "Receipt ID": receipt_id,
        "Date": _parse_receipt_date(result.ticket.date),
        "Store": result.ticket.store_name,
        "Gross Total (€)": result.ticket.gross_total,
        "Discounts (€)": result.ticket.discount_total,
        "Total Paid (€)": result.ticket.total_paid,
        "Line Count": len(result.products),
        "Payment Method": result.ticket.payment_method,
        "Notes": _build_receipt_notes(result),
    }

    _write_row(worksheet=worksheet, header_map=header_map, values=values, id_header="Receipt ID")


def _append_price_history_rows(
    worksheet,
    header_map: dict[str, int],
    receipt_id: str,
    record_ids: list[str],
    accepted_products: list[dict[str, Any]],
    result: ReceiptExtraction,
) -> None:
    """
    Append accepted product observations to Price_History.
    """

    if len(record_ids) != len(accepted_products):
        raise ValueError("The number of Record IDs does not match the number of accepted products.")

    receipt_date = _parse_receipt_date(result.ticket.date)

    store_name = (result.ticket.store_name)

    for record_id, product in zip(
        record_ids,
        accepted_products,
        strict=True, # Use strict=True to ensure that the lengths of record_ids and accepted_products match exactly
    ): # for each accepted product, append a row to the Price_History worksheet with the corresponding Record ID and receipt-level information
        notes = (
            f"Match: {product['match_method']}; "
            f"confidence: {product['confidence']:.2f}"
        )

        values = {
            "Record ID": record_id,
            "Receipt ID": receipt_id,
            "Date": receipt_date,
            "Store": store_name,
            "Product ID": product["product_id"],
            "Receipt/Original Product": product["receipt_name"],
            "Original Quantity": product["quantity_original"],
            "Comparable Quantity": product["comparable_quantity"],
            "Comparable Unit": product["comparable_unit"],
            "Gross Price (€)": product["gross_price"],
            "Discount (€)": product["discount_amount"],
            "Price Paid (€)": product["price_paid"],
            "Comparable Price": product["comparable_price"],
            "Notes": notes,
        }

        _write_row(worksheet=worksheet, header_map=header_map, values=values, id_header="Record ID")


def _append_review_rows(
    worksheet,
    header_map: dict[str, int],
    receipt_id: str,
    review_ids: list[str],
    review_items: list[dict[str, Any]],
) -> None:
    """
    Append review items to the Review worksheet.
    """

    if len(review_ids) != len(
        review_items
    ):
        raise ValueError(
            "The number of Review IDs does not match "
            "the number of review items."
        )

    for review_id, item in zip(
        review_ids,
        review_items,
        strict=True,
    ):
        values = {
            "Review ID": review_id,
            "Receipt ID": receipt_id,
            "Original Text": item["original_text"],
            "Issue Type": item["issue_type"],
            "AI Suggestion": item["ai_suggestion"],
            "Confidence": item["confidence"],
            "Status": item["status"],
            "Correction": item["correction"],
        }

        _write_row(worksheet=worksheet, header_map=header_map, values=values, id_header="Review ID")


def write_processed_receipt(
    source_excel_path: str | Path,
    output_excel_path: str | Path,
    result: ReceiptExtraction,
    accepted_products: list[dict[str, Any]],
    review_items: list[dict[str, Any]],
    receipt_id: str,
    record_ids: list[str],
    review_ids: list[str],
) -> None:
    """
    Write one processed receipt to a new workbook file.

    The source workbook is not overwritten.
    """

    source_excel_path = Path(source_excel_path)

    output_excel_path = Path(output_excel_path)

    if not source_excel_path.exists():
        raise FileNotFoundError(
            f"Source Excel file not found: {source_excel_path}"
        )

    output_excel_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(filename=source_excel_path)

    try:
        required_sheets = {RECEIPTS_SHEET, PRICE_HISTORY_SHEET, REVIEW_SHEET}

        missing_sheets = required_sheets - set(workbook.sheetnames) 

        if missing_sheets: # If missing_sheets is not empty, raise an error with the missing sheet names
            missing = ", ".join(sorted(missing_sheets))

            raise ValueError(
                f"Workbook is missing required sheets: {missing}"
            )

        receipts_sheet = workbook[RECEIPTS_SHEET]
        price_history_sheet = workbook[PRICE_HISTORY_SHEET]
        review_sheet = workbook[REVIEW_SHEET]

        receipts_header_map = _build_header_map(receipts_sheet)
        price_history_header_map = _build_header_map(price_history_sheet)
        review_header_map = _build_header_map(review_sheet)

        _validate_headers(
            header_map=receipts_header_map,
            required_headers=RECEIPTS_REQUIRED_HEADERS,
            sheet_name=RECEIPTS_SHEET,
        )

        _validate_headers(
            header_map=price_history_header_map,
            required_headers=PRICE_HISTORY_REQUIRED_HEADERS,
            sheet_name=PRICE_HISTORY_SHEET,
        )

        _validate_headers(
            header_map=review_header_map,
            required_headers=REVIEW_REQUIRED_HEADERS,
            sheet_name=REVIEW_SHEET,
        )

        _append_receipt(
            worksheet=receipts_sheet,
            header_map=receipts_header_map,
            receipt_id=receipt_id,
            result=result,
        )

        _append_price_history_rows(
            worksheet=price_history_sheet,
            header_map=price_history_header_map,
            receipt_id=receipt_id,
            record_ids=record_ids,
            accepted_products=accepted_products,
            result=result,
        )

        _append_review_rows(
            worksheet=review_sheet,
            header_map=review_header_map,
            receipt_id=receipt_id,
            review_ids=review_ids,
            review_items=review_items,
        )

        workbook.save(output_excel_path)

    finally:
        workbook.close()